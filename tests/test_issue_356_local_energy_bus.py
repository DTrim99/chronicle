"""Regression tests for the chronicle#356 local energy, area classification and bus facts.

- DESNZ sub-national electricity and gas for Northern Ireland, every year and measure.
- DESNZ's gas year as its own period type (gas_year), on the GB local-authority gas
  rows and the new packages.
- DESNZ domestic meters and consumption by LSOA and Scottish data zone, 2023 and 2024
  (the workbooks are R2-only and fetched by sha256 when absent).
- Raw-only registrations: the MSOA/intermediate-zone workbooks and the rural-urban
  lookups for England and Wales, Scotland, Northern Ireland and DfT's BUS01 scopes.
- DfT BUS04ai total operating revenue for Scotland, Wales and Great Britain.

Values are the publishers' as fetched on 2026-10-15.
"""

from __future__ import annotations

import re
import subprocess
import zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
import yaml

from chronicle.bundle import UK_BUNDLE_SOURCES
from chronicle.core import validate_facts
from chronicle.source_package import (
    SOURCE_PACKAGE_ALIASES,
    load_source_package,
    read_registered_artifact,
)
from chronicle.sources.cells import SourceArtifactMetadata, source_cells_from_xlsx

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = REPOSITORY_ROOT / "db" / "data"

NI_ELECTRICITY = "desnz-subnational-electricity-consumption-northern-ireland-2024"
NI_GAS = "desnz-subnational-gas-consumption-northern-ireland-2024"
GB_GAS = "desnz-subnational-gas-consumption-2024"
LSOA_GAS = "desnz-lsoa-domestic-gas-consumption-2024"
LSOA_ELECTRICITY = "desnz-lsoa-domestic-electricity-consumption-2024"
BUS04 = "dft-bus04i-operating-revenue-2025"
NI = "N92000002"

RAW_ONLY = {
    "desnz/msoa_domestic_gas_consumption_2010_2024": "74515cdfe094ceddf9de3273440ba18f9abc064fff4e055a3265144e706c13a1",
    "desnz/msoa_domestic_electricity_consumption_2010_2024": "9b880e6985d90fa7dbab7e67041fa0dc32948193f2dedcf2459ef59b99b0a40a",
    "ons/ruc2021_lookups_march_2026": "bb19be3e628da4351887863ac310145ee476025438d982aeffc2443140454b36",
    "scotgov/urban_rural_classification_2022_dz2022": "e70e5660f16302e72f6c364ccc03548de30c8c3a1f399e900e5886b7d4c51024",
    "scotgov/urban_rural_classification_2022_oa2022": "b512c345d540dfa0f492a1a02ec45f54857c720455c2189cedb3070b73ca6dc7",
    "nisra/settlement_2015_urban_rural_lookup": "1f9e4c3d2d1ccadbf80eb52f8dc6b8f8867ca0a1aec327925df03cb993ac0ff0",
    "nisra/data_zone_2021_lookups_v3": "8e2e1f6daaddd2f5b6887ccab1bf9d0bf179f6d0d2ba60ef8e0ebb495a3e1999",
    "defra/ruc2011_local_authority_districts_england": "26115b18425b99f69693627ffe26917ac804937cde0e5d2523b0791e771a073e",
    "ons/ruc2011_counties_england": "3224816ca18638dd86849b60b90359b7a54ba9f1c472bbc04db4feabea2fb9a2",
}


@pytest.mark.parametrize(
    "alias", [NI_ELECTRICITY, NI_GAS, LSOA_GAS, LSOA_ELECTRICITY, BUS04]
)
def test_packages_are_registered(alias):
    assert alias in SOURCE_PACKAGE_ALIASES
    assert alias in UK_BUNDLE_SOURCES


@pytest.fixture(scope="module")
def ni_electricity():
    return load_source_package(NI_ELECTRICITY).build_facts(2025)


@pytest.fixture(scope="module")
def ni_gas():
    return load_source_package(NI_GAS).build_facts(2025)


def test_ni_electricity_years_periods_and_geographies(ni_electricity):
    assert validate_facts(ni_electricity).valid
    assert len(ni_electricity) == 4616
    periods = {(f.period.type, f.period.value) for f in ni_electricity}
    assert {v for t, v in periods if t == "calendar_year"} == {2009, 2010, 2011}
    assert {v for t, v in periods if t == "fiscal_year"} == set(range(2012, 2025))
    fy2024 = next(f for f in ni_electricity if f.period.value == 2024)
    assert (fy2024.period_coverage.start_date, fy2024.period_coverage.end_date) == (
        "2024-04-01",
        "2025-03-31",
    )
    districts = {f.geography.id for f in ni_electricity if f.period.value == 2024}
    assert {d for d in districts if d.startswith("N09")} == {
        f"N090000{i:02d}" for i in range(1, 12)
    }
    assert "desnz:northern_ireland_unallocated" in districts
    old = {f.geography.id for f in ni_electricity if f.period.value == 2014}
    assert "UKN0301" in old and len([d for d in old if d.startswith("UKN0")]) == 26


def test_ni_electricity_all_domestic_matches_the_gb_concepts(ni_electricity):
    latest = {
        (f.filters["domestic_scope"], f.filters["metric"]): f
        for f in ni_electricity
        if f.geography.id == NI and f.period.value == 2024
    }
    meters = latest[("all_domestic", "meter_count")]
    consumption = latest[("all_domestic", "consumption")]
    assert meters.value == 848_085
    assert consumption.value == pytest.approx(2_695_722_393.297, rel=1e-9)
    assert meters.layout.measure_id == "domestic_meter_count"
    assert (
        meters.measure.concept == "desnz.subnational.domestic_electricity_meter_count"
    )
    assert (
        consumption.measure.concept
        == "desnz.subnational.domestic_electricity_consumption"
    )
    gb = load_source_package(
        "desnz-subnational-electricity-consumption-2024"
    ).build_facts(2026)
    gb_keys = {
        (f.measure.concept, tuple(sorted(f.filters.items())), f.layout.measure_id)
        for f in gb
        if f.geography.id == "K03000001"
    }
    assert (
        meters.measure.concept,
        tuple(sorted(meters.filters.items())),
        meters.layout.measure_id,
    ) in gb_keys
    assert latest[("all_domestic", "median_consumption_per_meter")].value == 2711
    assert latest[("all_non_domestic", "meter_count")].measure.concept == (
        "desnz.subnational.non_domestic_electricity_meter_count"
    )


def test_ni_gas_years_are_dated_and_suppressed_cells_emit_nothing(ni_gas):
    assert validate_facts(ni_gas).valid
    assert {f.period.type for f in ni_gas} == {"gas_year"}
    dates = {
        f.period.value: (f.period_coverage.start_date, f.period_coverage.end_date)
        for f in ni_gas
    }
    assert dates[2015] == ("2014-10-01", "2015-09-30")
    assert dates[2016] == ("2016-07-15", "2017-07-15")
    assert dates[2024] == ("2024-05-15", "2025-05-15")
    assert not [f for f in ni_gas if "median" in f.filters["metric"]]
    assert not [f for f in ni_gas if "unallocated" in f.geography.id]
    # '[no data]' in Fermanagh and Omagh's and Mid Ulster's non-domestic and all-meter
    # cells from 2019; their domestic cells are ported.
    for district in ("N09000006", "N09000009"):
        scopes = {
            f.filters["domestic_scope"]
            for f in ni_gas
            if f.geography.id == district and f.period.value == 2024
        }
        assert scopes == {"all_domestic"}
    latest = {
        (f.filters["domestic_scope"], f.filters["metric"]): f.value
        for f in ni_gas
        if f.geography.id == NI and f.period.value == 2024
    }
    assert latest[("all_domestic", "meter_count")] == 322_974
    assert latest[("all_domestic", "consumption")] == pytest.approx(
        2_984_987_252.54, rel=1e-9
    )


def test_gb_gas_rows_are_gas_years_with_unchanged_values():
    facts = load_source_package(GB_GAS).build_facts(2026)
    assert {f.period.type for f in facts} == {"gas_year"}
    assert {f.period.value for f in facts} == {2023, 2024}
    # Great Britain's gas years publish only mid-May endpoints: no dates are stamped.
    assert {
        (f.period_coverage.start_date, f.period_coverage.end_date) for f in facts
    } == {(None, None)}
    labels = {f.period_coverage.source_period_label for f in facts}
    assert labels == {
        "2023 (mid-May 2023 to mid-May 2024)",
        "2024 (mid-May 2024 to mid-May 2025)",
    }
    assert {f.layout.record_set_id for f in facts} == {
        "desnz.subnational.gas.cy2023",
        "desnz.subnational.gas.cy2024",
    }


def _skip_unless_available(alias: str) -> None:
    """Skip, as chronicle#343's raw-byte tests do, when an R2-only workbook is neither
    committed nor cached and fetching (LEDGER_SOURCE_ARTIFACT_FETCH, which CI sets) is off.
    """
    directory = {
        LSOA_GAS: "desnz/lsoa_domestic_gas_consumption_2024",
        LSOA_ELECTRICITY: "desnz/lsoa_domestic_electricity_consumption_2024",
    }[alias]
    try:
        read_registered_artifact(DATA_ROOT / directory / "manifest.yaml", 2025)
    except FileNotFoundError as error:
        pytest.skip(f"R2-only workbook unavailable: {error}")


@pytest.mark.parametrize(
    "alias,fuel,kept",
    [
        (LSOA_GAS, "gas", {2023: 41_457, 2024: 41_460}),
        (LSOA_ELECTRICITY, "electricity", {2023: 43_060, 2024: 43_063}),
    ],
)
def test_lsoa_meters_and_consumption(alias, fuel, kept):
    _skip_unless_available(alias)
    package = load_source_package(alias)
    assert package.artifact.parser == "xlsx_read_only_used_range"
    assert package.artifact.sheets == ("2023", "2024")
    facts = package.build_facts(2025)
    assert validate_facts(facts).valid
    rows = Counter(f.period.value for f in facts)
    assert rows == {year: 2 * count for year, count in kept.items()}
    assert {f.filters["metric"] for f in facts} == {"meter_count", "consumption"}
    assert {f.filters["fuel"] for f in facts} == {fuel}
    vintages = Counter(
        f.geography.geography_vintage
        if hasattr(f.geography, "geography_vintage")
        else f.geography.vintage
        for f in facts
        if f.period.value == 2024
    )
    assert set(vintages) == {"lsoa_2021", "dz_2022", "current"}
    assert {f.geography.id for f in facts if f.geography.vintage == "current"} == {
        "desnz:great_britain_unallocated"
    }
    assert {f.period.type for f in facts} == (
        {"gas_year"} if fuel == "gas" else {"calendar_year"}
    )
    first = next(
        f
        for f in facts
        if f.geography.id == "E01011954"
        and f.period.value == 2024
        and f.filters["metric"] == "meter_count"
    )
    assert first.value == (1014 if fuel == "gas" else 1015)
    assert first.geography.name == "Hartlepool 001A"


@pytest.mark.parametrize(
    "directory",
    [
        "desnz/lsoa_domestic_gas_consumption_2024",
        "desnz/lsoa_domestic_electricity_consumption_2024",
    ],
)
def test_lsoa_workbooks_are_r2_only_and_git_ignored(directory):
    manifest = yaml.safe_load((DATA_ROOT / directory / "manifest.yaml").read_text())
    (entry,) = manifest["files"].values()
    assert entry["storage"]["local"] is False
    assert entry["storage"]["r2"]["key"].startswith("raw/uk/desnz/desnz-lsoa-domestic-")
    ignored = subprocess.run(
        [
            "git",
            "check-ignore",
            "-q",
            str(Path("db/data") / directory / entry["filename"]),
        ],
        cwd=REPOSITORY_ROOT,
    )
    assert ignored.returncode == 0


def test_read_only_parse_matches_the_full_parse():
    path = (
        DATA_ROOT
        / "desnz/subnational_gas_consumption_northern_ireland_2024"
        / "Northern_Ireland_subnational_gas_consumption_2015-2024.xlsx"
    )
    artifact = SourceArtifactMetadata(
        source_name="desnz",
        source_table="t",
        source_file=path.name,
        url=None,
        vintage="v",
        sha256="0",
        size_bytes=0,
        extracted_at="2026-10-15",
        extraction_method="test",
    )
    sheets = ("2023", "2024")
    full = source_cells_from_xlsx(path.read_bytes(), artifact, sheets=sheets)
    streamed = source_cells_from_xlsx(
        path.read_bytes(), artifact, sheets=sheets, read_only=True
    )
    assert [
        (c.sheet_name, c.address, c.cell_type, c.raw_value, c.formula) for c in streamed
    ] == [(c.sheet_name, c.address, c.cell_type, c.raw_value, c.formula) for c in full]
    with pytest.raises(ValueError, match="must name the sheets"):
        source_cells_from_xlsx(path.read_bytes(), artifact, read_only=True)


def _workbook_with_understated_dimension() -> bytes:
    """A 5 x 3 sheet with a formula whose <dimension> claims only A1:B2."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "data"
    for row in range(1, 6):
        for column in range(1, 3):
            sheet.cell(row=row, column=column, value=row * 10 + column)
    for row in range(1, 6):
        sheet.cell(row=row, column=3, value=f"=A{row}+B{row}")
    sheet["A4"] = None
    written = BytesIO()
    workbook.save(written)
    rewritten = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(written.getvalue())) as source,
        zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data, replaced = re.subn(
                    rb'<dimension ref="A1:C5"\s*/>', b'<dimension ref="A1:B2"/>', data
                )
                assert replaced == 1
            target.writestr(item, data)
    return rewritten.getvalue()


def test_read_only_parse_measures_a_sheet_whose_dimension_understates_it():
    content = _workbook_with_understated_dimension()
    artifact = SourceArtifactMetadata(
        source_name="test",
        source_table="t",
        source_file="understated.xlsx",
        url=None,
        vintage="v",
        sha256="0",
        size_bytes=len(content),
        extracted_at="2026-10-15",
        extraction_method="test",
    )
    full = source_cells_from_xlsx(content, artifact, sheets=("data",))
    streamed = source_cells_from_xlsx(
        content, artifact, sheets=("data",), read_only=True
    )

    def key(cell):
        return (cell.address, cell.cell_type, cell.raw_value, cell.formula)

    assert len(full) == 15
    assert [key(c) for c in streamed] == [key(c) for c in full]
    # The formula column is captured in both parses, not only the values.
    assert [c.formula for c in streamed if c.address.startswith("C")] == [
        f"=A{row}+B{row}" for row in range(1, 6)
    ]


@pytest.mark.parametrize("directory,sha256", sorted(RAW_ONLY.items()))
def test_raw_only_registrations(directory, sha256):
    manifest = yaml.safe_load((DATA_ROOT / directory / "manifest.yaml").read_text())
    assert manifest["registration_kind"] == "raw_only"
    assert manifest["provenance_notes"]
    (entry,) = manifest["files"].values()
    assert entry["sha256"] == sha256
    publisher = directory.split("/")[0]
    assert entry["storage"]["r2"]["key"].startswith(f"raw/uk/{publisher}/")
    # The 9-15 MB workbooks no build parses are R2-only; the rest are committed.
    r2_only = entry["storage"].get("local") is False
    assert r2_only == (entry["size_bytes"] > 9_000_000)
    assert r2_only or (DATA_ROOT / directory / entry["filename"]).exists()
    package_dirs = {str(path) for path in SOURCE_PACKAGE_ALIASES.values()}
    assert directory not in package_dirs


def test_bus04_country_operating_revenue():
    facts = load_source_package(BUS04).build_facts(2025)
    values = {(f.geography.id, f.period.value): f.value for f in facts}
    assert set(values) == {
        (code, year)
        for code in ("S92000003", "W92000004", "K03000001")
        for year in (2024, 2025)
    }
    # DfT's single-year label is kept: 2024 is the year ending March 2024.
    assert values[("W92000004", 2024)] == pytest.approx(172_791_124.68, rel=1e-9)
    assert values[("W92000004", 2025)] == pytest.approx(203_655_522.23, rel=1e-9)
    assert values[("K03000001", 2025)] == pytest.approx(7_727_416_075.13, rel=1e-9)
    assert {f.measure.concept for f in facts} == {
        "dft.local_bus_total_estimated_operating_revenue"
    }
    # England is BUS05ai's 'Total estimated operating revenue', already in the feed.
    england = [
        f
        for f in load_source_package("dft-bus05i-revenue-support-2025").build_facts(
            2025
        )
        if f.geography.id == "E92000001"
        and f.measure.concept == "dft.local_bus_total_estimated_operating_revenue"
        and f.period.value == 2025
    ]
    assert len(england) == 1 and england[0].value == pytest.approx(
        6_616_074_977.28, rel=1e-9
    )
